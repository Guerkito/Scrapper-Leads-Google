import io
import os
import sys
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import db
from ui.crm import (
    _build_professional_excel,
    _build_portfolio_zip,
    _crm_table_key,
    _prepare_excel_export,
    _split_portfolio_batches,
    _validate_export_identity,
)


def _sample_leads():
    return pd.DataFrame([
        {
            "id": 101,
            "place_id": "0x1:0x101",
            "nombre": "Café Uno",
            "ciudad": "Bogotá",
            "departamento": "Cundinamarca",
            "pais": "Colombia",
            "nicho": "Cafeterías",
            "sector": "Gastronomía",
            "tipo": "B2C",
            "telefono": "300 111 2233",
            "telefono_e164": "+573001112233",
            "maps_url": "https://www.google.com/maps/place/Cafe+Uno/data=!1s0x1:0x101",
            "rating": 3.5,
            "reseñas": 5,
            "tiene_web": False,
            "calificacion": "frio",
            "estado": "Nuevo",
            "estado_contacto": "sin_contactar",
        },
        {
            "id": 202,
            "place_id": "0x2:0x202",
            "nombre": "Café Dos",
            "ciudad": "Bogotá",
            "departamento": "Cundinamarca",
            "pais": "Colombia",
            "nicho": "Cafeterías",
            "sector": "Gastronomía",
            "tipo": "B2C",
            "telefono": "300 111 2233",
            "telefono_e164": "+573001112233",
            "maps_url": "https://www.google.com/maps/place/Cafe+Dos/data=!1s0x2:0x202",
            "rating": 4.9,
            "reseñas": 200,
            "tiene_web": False,
            "calificacion": "oro",
            "estado": "Nuevo",
            "estado_contacto": "sin_contactar",
        },
        {
            "id": 303,
            "place_id": "0x3:0x303",
            "nombre": "Café Tres",
            "ciudad": "Bogotá",
            "departamento": "Cundinamarca",
            "pais": "Colombia",
            "nicho": "Cafeterías",
            "sector": "Gastronomía",
            "tipo": "B2C",
            "telefono": "301 222 3344",
            "telefono_e164": "+573012223344",
            "maps_url": "https://www.google.com/maps/place/Cafe+Tres/data=!1s0x3:0x303",
            "rating": 4.8,
            "reseñas": 100,
            "tiene_web": True,
            "calificacion": "bueno",
            "estado": "Nuevo",
            "estado_contacto": "no_contactar",
        },
    ])


def test_excel_preparation_keeps_identity_after_sorting():
    source = _sample_leads()
    exported = _prepare_excel_export(source)

    assert exported["Lead ID"].tolist()[0] == 202
    for _, row in exported.iterrows():
        original = source[source["id"] == row["Lead ID"]].iloc[0]
        assert row["Empresa"] == original["nombre"]
        assert row["Place ID"] == original["place_id"]
        assert row["Maps"] == original["maps_url"]
        assert str(row["Telefono E164"]).removeprefix("'") == original["telefono_e164"]


def test_export_identity_validator_blocks_crossed_url():
    source = _sample_leads()
    exported = _prepare_excel_export(source).copy()
    exported["Telefono E164"] = exported["Telefono E164"].map(
        lambda value: str(value).removeprefix("'")
    )
    _validate_export_identity(source, exported)

    exported.loc[exported["Lead ID"] == 101, "Maps"] = source.loc[
        source["id"] == 202, "maps_url"
    ].iloc[0]
    with pytest.raises(ValueError, match="lead 101"):
        _validate_export_identity(source, exported)


def test_workbook_links_and_call_queue_stay_bound_to_lead_id():
    source = _sample_leads()
    workbook = load_workbook(io.BytesIO(_build_professional_excel(source)))

    assert "Lista para llamar" in workbook.sheetnames
    leads = workbook["Leads"]
    headers = {cell.value: cell.column for cell in leads[1]}
    rows_by_id = {
        leads.cell(row=row, column=headers["Lead ID"]).value: row
        for row in range(2, leads.max_row + 1)
    }
    for _, original in source.iterrows():
        row = rows_by_id[original["id"]]
        maps_cell = leads.cell(row=row, column=headers["Maps"])
        phone_cell = leads.cell(row=row, column=headers["Telefono E164"])
        assert maps_cell.value == "Abrir Maps"
        assert maps_cell.hyperlink.target == original["maps_url"]
        assert phone_cell.value == original["telefono_e164"]

    queue = workbook["Lista para llamar"]
    queue_headers = {cell.value: cell.column for cell in queue[1]}
    queue_ids = [
        queue.cell(row=row, column=queue_headers["Lead ID"]).value
        for row in range(2, queue.max_row + 1)
    ]
    assert queue_ids == [202]
    queue_maps = queue.cell(row=2, column=queue_headers["Maps"])
    assert queue_maps.hyperlink.target == source.loc[
        source["id"] == 202, "maps_url"
    ].iloc[0]


def test_crm_table_key_changes_with_row_order():
    assert _crm_table_key([1, 2, 3]) == _crm_table_key([1, 2, 3])
    assert _crm_table_key([1, 2, 3]) != _crm_table_key([3, 2, 1])


def test_portfolio_batches_keep_divi_watson_onyx_and_unassigned_separate():
    source = _sample_leads()
    source["product_keys"] = [
        "divi_restaurantes",
        "watson_clinic,onyx_web",
        "campana_anterior",
    ]
    source["productos_objetivo"] = [
        "DIVI · Restaurantes",
        "Watson Clinic,ONYX · Páginas web",
        "",
    ]

    batches = _split_portfolio_batches(source)

    assert batches["divi"]["id"].tolist() == [101]
    assert batches["watson"]["id"].tolist() == [202]
    assert batches["onyx"]["id"].tolist() == [202]
    assert batches["unassigned"]["id"].tolist() == [303]
    assert batches["watson"]["lote_exportado"].tolist() == ["Watson"]


def test_batch_zip_writes_one_csv_per_selected_portfolio_without_mixing_rows():
    source = _sample_leads()
    source["product_keys"] = ["divi_restaurantes", "watson_clinic", "onyx_web"]
    source["productos_objetivo"] = [
        "DIVI · Restaurantes", "Watson Clinic", "ONYX · Páginas web"
    ]
    batches = _split_portfolio_batches(source)

    package = _build_portfolio_zip(
        batches, ["divi", "watson"], file_format="csv", stamp="20260724_1200"
    )

    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        assert archive.namelist() == [
            "DIVI_Leads_20260724_1200.csv",
            "WATSON_Leads_20260724_1200.csv",
        ]
        divi = pd.read_csv(
            io.BytesIO(archive.read(archive.namelist()[0])), sep=";", encoding="utf-8-sig"
        )
        watson = pd.read_csv(
            io.BytesIO(archive.read(archive.namelist()[1])), sep=";", encoding="utf-8-sig"
        )

    assert divi["Lead ID"].tolist() == [101]
    assert watson["Lead ID"].tolist() == [202]
    assert divi["Lote exportado"].tolist() == ["DIVI"]
    assert watson["Lote exportado"].tolist() == ["Watson"]


def _insert_lead_with_opportunities(conn, name, opportunities):
    lead_id = conn.execute(
        "INSERT INTO leads(nombre, ciudad, pais) VALUES (?, 'Bogotá', 'Colombia')",
        (name,),
    ).lastrowid
    for product_key, product_label in opportunities:
        conn.execute(
            """
            INSERT INTO lead_opportunities
                (lead_id, product_key, product_label, segment_key)
            VALUES (?, ?, ?, '')
            """,
            (lead_id, product_key, product_label),
        )
    return lead_id


def test_delete_portfolio_batch_preserves_lead_shared_with_another_portfolio(
    monkeypatch, tmp_path
):
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "delete-batch.db"))
    db.init_db()
    with db.open_conn() as conn:
        divi_only = _insert_lead_with_opportunities(
            conn, "Solo DIVI", [("divi_restaurantes", "DIVI · Restaurantes")]
        )
        shared = _insert_lead_with_opportunities(
            conn,
            "Compartido",
            [
                ("divi_restaurantes", "DIVI · Restaurantes"),
                ("watson_clinic", "Watson Clinic"),
            ],
        )
        watson_only = _insert_lead_with_opportunities(
            conn, "Solo Watson", [("watson_clinic", "Watson Clinic")]
        )

    invalidations = []
    import services.leads as lead_service
    monkeypatch.setattr(
        lead_service, "invalidate_leads_cache", lambda: invalidations.append(True)
    )

    result = db.delete_portfolio_batch(
        [divi_only, shared, watson_only], "divi"
    )

    assert result == {
        "requested": 3,
        "matched": 2,
        "leads_deleted": 1,
        "leads_preserved": 1,
        "associations_deleted": 2,
    }
    assert invalidations == [True]
    with db.open_conn() as conn:
        remaining_leads = {
            row[0] for row in conn.execute("SELECT id FROM leads").fetchall()
        }
        remaining_opportunities = conn.execute(
            "SELECT lead_id, product_key FROM lead_opportunities ORDER BY lead_id"
        ).fetchall()
    assert remaining_leads == {shared, watson_only}
    assert remaining_opportunities == [
        (shared, "watson_clinic"),
        (watson_only, "watson_clinic"),
    ]


def test_delete_unassigned_batch_rechecks_membership_before_deleting(
    monkeypatch, tmp_path
):
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "delete-unassigned.db"))
    db.init_db()
    with db.open_conn() as conn:
        unknown = _insert_lead_with_opportunities(
            conn, "Campaña antigua", [("legacy_offer", "Oferta anterior")]
        )
        watson = _insert_lead_with_opportunities(
            conn, "Watson", [("watson_clinic", "Watson Clinic")]
        )

    result = db.delete_portfolio_batch([unknown, watson], "unassigned")

    assert result["matched"] == 1
    assert result["leads_deleted"] == 1
    with db.open_conn() as conn:
        remaining = conn.execute("SELECT id FROM leads").fetchall()
    assert remaining == [(watson,)]
